"""FastAPI backend for Office of CEO Insights Builder."""

import os
import sys
import json
import asyncio
from datetime import datetime
from typing import Optional
from contextlib import asynccontextmanager

# Force UTF-8 stdout/stderr so emoji/unicode in logs don't crash on Windows
# (Windows default 'charmap' codec can't encode characters like ✓, →, ⚠).
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv

from models import (
    Session,
    ChatMessage,
    TokenUsage,
    SessionTokenUsage,
    Collaborator,
    DraftDocument,
    DraftEvaluation,
    RubricCriterion,
    CreateSessionRequest,
    SendMessageRequest,
    EditMessageRequest,
    AddCollaboratorRequest,
    AssignRequest,
    SaveDraftRequest,
    EvaluateDraftRequest,
    RephraseRequest,
    TTSRequest,
    StreamEvent,
)
from agents import AgentOrchestrator
from fastapi.responses import StreamingResponse, Response
import io

load_dotenv()

# In-memory session store
sessions: dict[str, Session] = {}
orchestrator: Optional[AgentOrchestrator] = None


def _safe_filename(title: str) -> str:
    """Sanitize a session title into a safe filename stem."""
    import re
    stem = re.sub(r"[^\w\- ]", "", title or "document").strip().replace(" ", "_")
    return stem[:60] or "document"


def _markdown_to_docx(markdown: str, title: str) -> "io.BytesIO":
    """Convert markdown content into a .docx document (in-memory buffer).

    Handles headings (#..######), bullet lists (-/*), numbered lists,
    simple pipe tables, bold (**text**), and paragraphs.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor
    import re

    doc = Document()

    def add_runs_with_bold(paragraph, text):
        # Split on **bold** segments and add runs accordingly
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part:
                paragraph.add_run(part)

    lines = (markdown or "").split("\n")
    i = 0
    table_buffer = []

    def flush_table():
        nonlocal table_buffer
        if not table_buffer:
            return
        # Parse pipe table rows
        rows = []
        for raw in table_buffer:
            cells = [c.strip() for c in raw.strip().strip("|").split("|")]
            rows.append(cells)
        # Drop separator rows like |---|---|
        rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
        if rows:
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Light Grid Accent 1"
            for r_idx, r in enumerate(rows):
                cells = table.add_row().cells
                for c_idx in range(ncols):
                    val = r[c_idx] if c_idx < len(r) else ""
                    cells[c_idx].text = val.replace("**", "")
        table_buffer = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # Table detection (line contains pipes and looks like a row)
        if stripped.startswith("|") and stripped.count("|") >= 2:
            table_buffer.append(line)
            i += 1
            continue
        else:
            flush_table()

        if not stripped:
            i += 1
            continue

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).replace("**", "")
            if level == 1:
                h = doc.add_heading(text, level=0)
            else:
                h = doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^---+$", stripped):
            i += 1
            continue

        # Bullet list
        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, bullet_match.group(1))
            i += 1
            continue

        # Numbered list
        num_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if num_match:
            p = doc.add_paragraph(style="List Number")
            add_runs_with_bold(p, num_match.group(1))
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        add_runs_with_bold(p, stripped)
        i += 1

    flush_table()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


@asynccontextmanager
async def lifespan(app: FastAPI):
    """Initialize resources on startup."""
    global orchestrator
    orchestrator = AgentOrchestrator()
    try:
        if orchestrator.client:
            await orchestrator.initialize()
    except Exception as e:
        print(f"Warning: Agent initialization deferred - {e}")
    yield


app = FastAPI(
    title="Office of CEO - Insights Builder API",
    description="AI-powered executive insights for customer meetings",
    version="1.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# --- Session Management ---


@app.get("/api/sessions")
async def list_sessions():
    """List all sessions."""
    return [
        {
            "id": s.id,
            "title": s.title,
            "message_count": len(s.messages),
            "token_usage": s.token_usage.cumulative.model_dump(),
            "collaborator_count": len(s.collaborators),
            "assignee": s.assignee,
            "has_draft": bool(s.draft.content.strip()),
            "created_at": s.created_at.isoformat(),
            "updated_at": s.updated_at.isoformat(),
        }
        for s in sorted(sessions.values(), key=lambda x: x.updated_at, reverse=True)
    ]


@app.post("/api/sessions")
async def create_session(req: CreateSessionRequest):
    """Create a new session."""
    session = Session(title=req.title or "New Session")
    sessions[session.id] = session
    return {"id": session.id, "title": session.title}


@app.get("/api/sessions/{session_id}")
async def get_session(session_id: str):
    """Get session details with messages."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    return session.model_dump()


@app.delete("/api/sessions/{session_id}")
async def delete_session(session_id: str):
    """Delete a session."""
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")
    del sessions[session_id]
    return {"status": "deleted"}


@app.put("/api/sessions/{session_id}/title")
async def update_session_title(session_id: str, req: dict):
    """Update session title."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    session.title = req.get("title", session.title)
    session.updated_at = datetime.utcnow()
    return {"id": session.id, "title": session.title}


# --- Message Editing ---


@app.put("/api/messages/edit")
async def edit_message(req: EditMessageRequest):
    """Edit an existing message content."""
    session = sessions.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")

    for msg in session.messages:
        if msg.id == req.message_id:
            if not msg.is_edited:
                msg.original_content = msg.content
            msg.content = req.new_content
            msg.is_edited = True
            msg.edited_by = req.user_name
            session.updated_at = datetime.utcnow()
            return msg.model_dump()

    raise HTTPException(status_code=404, detail="Message not found")


# --- Collaboration ---


@app.post("/api/sessions/{session_id}/collaborators")
async def add_collaborator(session_id: str, req: AddCollaboratorRequest):
    """Invite a collaborator to a session."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    # Avoid duplicate emails
    for c in session.collaborators:
        if c.email.lower() == req.email.lower():
            return c.model_dump()
    collab = Collaborator(name=req.name, email=req.email, role=req.role)
    session.collaborators.append(collab)
    session.updated_at = datetime.utcnow()
    return collab.model_dump()


@app.delete("/api/sessions/{session_id}/collaborators/{collaborator_id}")
async def remove_collaborator(session_id: str, collaborator_id: str):
    """Remove a collaborator from a session."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    session.collaborators = [
        c for c in session.collaborators if c.id != collaborator_id
    ]
    if session.assignee == collaborator_id:
        session.assignee = None
    session.updated_at = datetime.utcnow()
    return {"status": "removed"}


@app.put("/api/sessions/{session_id}/assign")
async def assign_session(session_id: str, req: AssignRequest):
    """Assign session ownership to a collaborator (or clear with null)."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    if req.collaborator_id and not any(
        c.id == req.collaborator_id for c in session.collaborators
    ):
        raise HTTPException(status_code=400, detail="Collaborator not found in session")
    session.assignee = req.collaborator_id
    session.updated_at = datetime.utcnow()
    return {"assignee": session.assignee}


# --- Draft Document ---


@app.get("/api/sessions/{session_id}/draft")
async def get_draft(session_id: str):
    """Get the draft document for a session."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    return session.draft.model_dump()


@app.put("/api/sessions/{session_id}/draft")
async def save_draft(session_id: str, req: SaveDraftRequest):
    """Save/update the draft document for a session."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    session.draft = DraftDocument(
        content=req.content,
        updated_at=datetime.utcnow(),
        updated_by=req.user_name,
    )
    session.updated_at = datetime.utcnow()
    return session.draft.model_dump()


@app.get("/api/sessions/{session_id}/draft/export/markdown")
async def export_markdown(session_id: str):
    """Export the draft as a Markdown file."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    filename = _safe_filename(session.title) + ".md"
    return Response(
        content=session.draft.content,
        media_type="text/markdown",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/sessions/{session_id}/draft/export/docx")
async def export_docx(session_id: str):
    """Export the draft as a Word (.docx) file."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    buffer = _markdown_to_docx(session.draft.content, session.title)
    filename = _safe_filename(session.title) + ".docx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# --- Rubric Evaluation ---


@app.post("/api/sessions/{session_id}/evaluate")
async def evaluate_draft(session_id: str, req: EvaluateDraftRequest):
    """Evaluate the draft against an executive-briefing rubric."""
    session = sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    content = req.content if req.content is not None else session.draft.content
    if not content.strip():
        raise HTTPException(status_code=400, detail="Draft is empty — nothing to evaluate")

    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, orchestrator.evaluate_draft, content)

    session.evaluation = DraftEvaluation(
        overall_score=result.get("overall_score", 0.0),
        criteria=[RubricCriterion(**c) for c in result.get("criteria", [])],
        summary=result.get("summary", ""),
        strengths=result.get("strengths", []),
        improvements=result.get("improvements", []),
    )
    session.updated_at = datetime.utcnow()
    return session.evaluation.model_dump()


# --- AI Rephrase ---


@app.post("/api/rephrase")
async def rephrase(req: RephraseRequest):
    """Rephrase a snippet of selected text using the GPT model deployment."""
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="No text provided to rephrase")
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(
        None, orchestrator.rephrase_text, req.text, req.instruction or "", req.tone or ""
    )
    return result


# --- Text-to-Speech (voice read-out) ---


@app.post("/api/tts")
async def tts(req: TTSRequest):
    """Synthesize speech from text using the AUDIO_MODEL realtime deployment.

    Returns a WAV audio payload. On any failure (model not configured or
    synthesis error) responds 503 so the client can fall back to the browser
    SpeechSynthesis API.
    """
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="No text provided to read aloud")
    try:
        audio = await orchestrator.synthesize_speech(req.text, req.voice or "alloy")
    except Exception as e:
        raise HTTPException(status_code=503, detail=str(e))
    return Response(
        content=audio,
        media_type="audio/wav",
        headers={"Cache-Control": "no-store"},
    )


@app.post("/api/chat/stream")
async def stream_chat(req: SendMessageRequest):
    """Stream agent responses using SSE."""
    session = sessions.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")

    # Add user message
    user_message = ChatMessage(role="user", content=req.content)
    session.messages.append(user_message)
    session.updated_at = datetime.utcnow()

    # Auto-title from first message only when the user hasn't named the session
    if len(session.messages) == 1 and (not session.title or session.title == "New Session"):
        session.title = req.content[:50] + ("..." if len(req.content) > 50 else "")

    # Build conversation history for agents
    history = [{"role": m.role, "content": m.content} for m in session.messages[:-1]]

    async def event_generator():
        """Generate SSE events from agent orchestration."""
        context_message_id = None
        customer_data_message_id = None
        insights_message_id = None
        context_content = ""
        customer_data_content = ""
        insights_content = ""

        try:
            async for event in orchestrator.orchestrate(req.content, history):
                event_type = event["type"]

                if event_type == "agent_started":
                    agent = event["agent"]
                    # Create a placeholder message
                    msg = ChatMessage(
                        role="assistant", content="", agent=agent
                    )
                    session.messages.append(msg)
                    if agent == "context-builder":
                        context_message_id = msg.id
                    elif agent == "customer-data":
                        customer_data_message_id = msg.id
                    else:
                        insights_message_id = msg.id

                    sse_data = json.dumps(
                        {
                            "type": "agent_started",
                            "agent": agent,
                            "message_id": msg.id,
                        }
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "token":
                    agent = event["agent"]
                    content = event["content"]

                    # Append to the appropriate message
                    if agent == "context-builder":
                        context_content += content
                        for m in session.messages:
                            if m.id == context_message_id:
                                m.content = context_content
                                break
                        msg_id = context_message_id
                    elif agent == "customer-data":
                        customer_data_content += content
                        for m in session.messages:
                            if m.id == customer_data_message_id:
                                m.content = customer_data_content
                                break
                        msg_id = customer_data_message_id
                    else:
                        insights_content += content
                        for m in session.messages:
                            if m.id == insights_message_id:
                                m.content = insights_content
                                break
                        msg_id = insights_message_id

                    sse_data = json.dumps(
                        {
                            "type": "token",
                            "agent": agent,
                            "content": content,
                            "message_id": msg_id,
                        }
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "agent_completed":
                    agent = event["agent"]
                    usage_data = event.get("usage", {})
                    usage = TokenUsage(**usage_data)

                    # Update token tracking (customer-data shares context_builder bucket)
                    if agent in ("context-builder", "customer-data"):
                        session.token_usage.context_builder.prompt_tokens += (
                            usage.prompt_tokens
                        )
                        session.token_usage.context_builder.completion_tokens += (
                            usage.completion_tokens
                        )
                        session.token_usage.context_builder.total_tokens += (
                            usage.total_tokens
                        )
                    else:
                        session.token_usage.insights.prompt_tokens += (
                            usage.prompt_tokens
                        )
                        session.token_usage.insights.completion_tokens += (
                            usage.completion_tokens
                        )
                        session.token_usage.insights.total_tokens += usage.total_tokens

                    # Update cumulative
                    session.token_usage.cumulative.prompt_tokens += usage.prompt_tokens
                    session.token_usage.cumulative.completion_tokens += (
                        usage.completion_tokens
                    )
                    session.token_usage.cumulative.total_tokens += usage.total_tokens

                    sse_data = json.dumps(
                        {
                            "type": "agent_completed",
                            "agent": agent,
                            "usage": usage.model_dump(),
                            "cumulative_usage": session.token_usage.cumulative.model_dump(),
                        }
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "context_graph":
                    session.graph_data = event.get("data")
                    sse_data = json.dumps(
                        {"type": "context_graph", "data": event.get("data")}
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "watermelon_data":
                    session.watermelon_data = event.get("data")
                    sse_data = json.dumps(
                        {"type": "watermelon_data", "data": event.get("data")}
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "scorecard_data":
                    session.scorecard_data = event.get("data")
                    sse_data = json.dumps(
                        {"type": "scorecard_data", "data": event.get("data")}
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "done":
                    sse_data = json.dumps(
                        {
                            "type": "done",
                            "cumulative_usage": session.token_usage.cumulative.model_dump(),
                        }
                    )
                    yield f"data: {sse_data}\n\n"

                elif event_type == "error":
                    sse_data = json.dumps(
                        {"type": "error", "error": event.get("error", "Unknown error")}
                    )
                    yield f"data: {sse_data}\n\n"

        except Exception as e:
            sse_data = json.dumps({"type": "error", "error": str(e)})
            yield f"data: {sse_data}\n\n"

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


# --- Health Check ---


@app.get("/api/health")
async def health_check():
    return {
        "status": "healthy",
        "agents": {
            "context_builder": orchestrator.client is not None if orchestrator else False,
            "insights": orchestrator.client is not None if orchestrator else False,
        },
    }


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
